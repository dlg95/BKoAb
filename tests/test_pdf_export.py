from io import BytesIO

from docx import Document
from pypdf import PdfReader

from bkoab.services.pdf_export import docx_bytes_to_pdf, find_soffice, merge_pdf_documents


def _minimal_pdf(page_label: str = "1") -> bytes:
    return f"""%PDF-1.4
1 0 obj<< /Type /Catalog /Pages 2 0 R >>endobj
2 0 obj<< /Type /Pages /Kids [3 0 R] /Count 1 >>endobj
3 0 obj<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] /Contents 4 0 R >>endobj
4 0 obj<< /Length 44 >>stream
BT /F1 12 Tf 20 100 Td ({page_label}) Tj ET
endstream endobj
xref
0 5
0000000000 65535 f 
0000000009 00000 n 
0000000058 00000 n 
0000000115 00000 n 
0000000214 00000 n 
trailer<< /Size 5 /Root 1 0 R >>
startxref
310
%%EOF
""".encode()


def test_find_soffice_respects_env(tmp_path, monkeypatch):
    fake = tmp_path / "soffice"
    fake.write_text("#!/bin/sh\nexit 0\n")
    fake.chmod(0o755)
    monkeypatch.setenv("BKOAB_SOFFICE", str(fake))
    assert find_soffice() == str(fake)


def test_merge_pdf_documents_preserves_page_count():
    merged = merge_pdf_documents([_minimal_pdf("A"), _minimal_pdf("B")])
    reader = PdfReader(BytesIO(merged))
    assert len(reader.pages) == 2


def test_merge_pdf_documents_skips_empty_parts():
    merged = merge_pdf_documents([_minimal_pdf("only"), b""])
    reader = PdfReader(BytesIO(merged))
    assert len(reader.pages) == 1


def test_docx_bytes_to_pdf_via_dxpdf(monkeypatch):
    monkeypatch.setenv("BKOAB_PDF_ENGINE", "dxpdf")
    doc = Document()
    doc.add_heading("BKoAb PDF-Test", 0)
    doc.add_paragraph("Umlaut: Größe, Straße, Gebäude.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Position"
    table.cell(0, 1).text = "Betrag"
    table.cell(1, 0).text = "Heizung"
    table.cell(1, 1).text = "12,34 €"
    buf = BytesIO()
    doc.save(buf)
    pdf = docx_bytes_to_pdf(buf.getvalue())
    assert pdf.startswith(b"%PDF")
    reader = PdfReader(BytesIO(pdf))
    assert len(reader.pages) >= 1
